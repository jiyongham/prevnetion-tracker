#!/usr/bin/env python3
"""
스크린샷 용 더미 데이터로 엑셀 치환.

사용:
  python3 scripts/dummy_data.py --excel targets.xlsx --output data/targets-dummy.xlsx
  python3 scripts/dummy_data.py --all      # 모든 엑셀을 -dummy로 생성

.gitignore에 *-dummy.xlsx 를 추가해 커밋되지 않게 함.
로컬에서만 쓰고, 캡처할 때 -dummy 파일을 data/ 에 복사해서 띄우면 됨.
"""

import argparse
import re
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


# 더미 데이터 사전
DUMMY_NAMES = [
    "홍길동", "김영희", "이순신", "박상민", "최지민",
    "이지은", "정민준", "우진서", "조예은", "강현준",
]
DUMMY_TEAMS = [
    "운영1팀", "운영2팀", "인프라팀", "보안팀", "네트워크팀",
]
DUMMY_SYSTEMS = [
    "[IT,인프라] 주요서비스 API", "[IT,인프라] 결제 시스템",
    "[금융,플랫폼] 메인 DB", "[금융,플랫폼] 백업",
    "[CMS,콘텐츠] 웹서버", "[CMS,콘텐츠] 캐시",
    "[모바일,앱] 푸시 서버", "[모바일,앱] 분석",
]
DUMMY_HOSTS = [
    f"host-{chr(97+i%5)}{j:02d}" for i in range(50) for j in range(10)
]

# 컬럼명 규칙에 안 걸리는 값들 (자유 텍스트 컬럼 안에 조직명이 섞여 나옴).
# 예: capacity.xlsx '자산 구분' 컬럼의 '그룹공통(데이터센터팀)'.
# 컬럼 단위가 아니라 값 안의 부분 문자열로 치환해야 그 칸의 나머지 정보가 안 날아간다.
VALUE_REPLACEMENTS = {
    "데이터센터팀": "인프라운영팀",
    "데이터센터": "인프라운영",
}


def should_replace(col_name: str) -> bool:
    """이 컬럼을 더미 데이터로 치환할까?"""
    patterns = [
        r"담당자|owner|사용자|name|고객|요청자",
        r"시스템|시스템명|ci명|application|app|service",
        r"호스트|host|server|서버",
        r"ip|주소",
        r"team|팀|org|부서|조직",
        r"group|그룹",
    ]
    col_lower = col_name.lower().strip()
    return any(re.search(p, col_lower) for p in patterns)


def dummy_value(col_name: str, row_idx: int) -> str:
    """컬럼 이름에 맞는 더미 값을 돌려준다."""
    col_lower = col_name.lower().strip()

    # 담당자 열 -> "이름-팀"
    if re.search(r"담당자|owner|사용자", col_lower):
        name = DUMMY_NAMES[row_idx % len(DUMMY_NAMES)]
        team = DUMMY_TEAMS[(row_idx // len(DUMMY_NAMES)) % len(DUMMY_TEAMS)]
        return f"{name}-{team}"

    # 팀 / 조직
    if re.search(r"^team|팀$|조직|org|부서|group", col_lower):
        return DUMMY_TEAMS[row_idx % len(DUMMY_TEAMS)]

    # 시스템명
    if re.search(r"시스템|시스템명|ci명|application|app|service", col_lower):
        return DUMMY_SYSTEMS[row_idx % len(DUMMY_SYSTEMS)]

    # 호스트명
    if re.search(r"호스트|host|server|서버", col_lower):
        return DUMMY_HOSTS[row_idx % len(DUMMY_HOSTS)]

    # IP
    if re.search(r"^ip|ip주소|주소|ipaddress", col_lower):
        base = 10 + (row_idx // 256)
        mid = (row_idx // 1) % 256
        host = row_idx % 256
        return f"10.{base}.{mid}.{host}"

    return str(row_idx)  # fallback


def anonymize_excel(input_path: str, output_path: str) -> None:
    """엑셀 파일을 더미 데이터로 치환하고 저장."""
    wb = load_workbook(input_path)

    for sheet in wb.sheetnames:
        ws = wb[sheet]

        # 첫 행을 헤더로 인식
        headers = {}
        for col_idx, cell in enumerate(ws[1], start=1):
            headers[col_idx] = cell.value

        # 2행부터 데이터 치환
        for row_idx in range(2, ws.max_row + 1):
            for col_idx, header in headers.items():
                cell = ws.cell(row=row_idx, column=col_idx)

                if header and should_replace(str(header)):
                    cell.value = dummy_value(str(header), row_idx - 2)
                    continue

                # 컬럼 단위 치환 대상이 아니어도, 값 안에 조직명이 섞여 있으면 부분 치환한다
                # (예: '자산 구분' 컬럼의 '그룹공통(데이터센터팀)').
                if isinstance(cell.value, str):
                    for old, new in VALUE_REPLACEMENTS.items():
                        if old in cell.value:
                            cell.value = cell.value.replace(old, new)

    wb.save(output_path)
    print(f"✅ {input_path} → {output_path}")


def main():
    parser = argparse.ArgumentParser(
        description="스크린샷 용 더미 데이터 생성"
    )
    parser.add_argument("--excel", help="변환할 엑셀 파일 경로")
    parser.add_argument("--output", help="출력 경로 (기본: -dummy.xlsx)")
    parser.add_argument("--all", action="store_true", help="data/ 안 모든 xlsx 변환")

    args = parser.parse_args()

    if args.all:
        data_dir = Path("data")
        for xlsx_file in data_dir.glob("*.xlsx"):
            if "-dummy" in xlsx_file.name or ".bak." in xlsx_file.name:
                continue
            output = xlsx_file.parent / f"{xlsx_file.stem}-dummy.xlsx"
            anonymize_excel(str(xlsx_file), str(output))
    elif args.excel:
        output = args.output or (args.excel.replace(".xlsx", "-dummy.xlsx"))
        anonymize_excel(args.excel, output)
    else:
        parser.print_help()


if __name__ == "__main__":
    main()
